#!/usr/bin/env python
"""Generate Fajiu debt-transfer SMS upload Excels.

Supports:
1) 债转短信*.xlsx with sheets 一转 / 二转
2) Flat case sheets (全部订单 / 短信 *.xlsx) with 担保方/受让方/反担保/_主表错误

Example:
  python generate_sms.py --input "短信 0910.xlsx" --out-dir ./out --date 0911
  python generate_sms.py --input 债转短信0904.xlsx --out-dir ./out --date 0904 --dry-run
"""

from __future__ import annotations

import argparse
import shutil
import sys
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook

SCRIPT_DIR = Path(__file__).resolve().parent
if str(SCRIPT_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPT_DIR))

from common import (  # noqa: E402
    COL_CURRENT_HOLDER,
    COL_FUND,
    COL_ORDER_ID,
    COL_SECOND_TRANSFEREE,
    COL_THIRD_TRANSFEROR,
    DEFAULT_FUND,
    clean_party,
    excel_saveas,
    find_default_template,
    normalize_oid_key,
    output_filename,
    short_of,
    to_order_id,
)


def _header_index(headers: list) -> dict[str, int]:
    return {str(name): i for i, name in enumerate(headers) if name is not None}


def _dedupe(oids: list) -> list:
    seen: set = set()
    out = []
    for oid in oids:
        key = normalize_oid_key(oid)
        if not key or key in seen:
            continue
        seen.add(key)
        out.append(oid)
    return out


def read_transfer_groups(input_path: Path) -> tuple[list[tuple[str, str]], dict[tuple[str, str], list]]:
    """Return (ordered keys, groups[(from,to)] -> order_ids)."""
    wb = load_workbook(input_path, read_only=True, data_only=True)
    sheet_names = wb.sheetnames
    groups: dict[tuple[str, str], list] = defaultdict(list)
    ordered: list[tuple[str, str]] = []

    def add_pair(from_party: str, to_party: str, oid) -> None:
        if not from_party or not to_party or oid in (None, ""):
            return
        key = (from_party, to_party)
        if key not in ordered:
            ordered.append(key)
        groups[key].append(to_order_id(oid))

    # Mode A: 一转 / 二转
    if "一转" in sheet_names or "二转" in sheet_names:
        for sheet_name, mode in (("一转", "first"), ("二转", "second")):
            if sheet_name not in sheet_names:
                continue
            ws = wb[sheet_name]
            rows = ws.iter_rows(values_only=True)
            headers = next(rows, None)
            if not headers:
                continue
            idx = _header_index(list(headers))
            for required in ("订单号", "受让方名称", "担保方名称", "反担保主体"):
                if required not in idx:
                    raise ValueError(f"{sheet_name} 缺少列: {required}")
            for row in rows:
                if not row or row[idx["订单号"]] in (None, ""):
                    continue
                oid = row[idx["订单号"]]
                gua = clean_party(row[idx["担保方名称"]])
                to = clean_party(row[idx["受让方名称"]])
                fan = clean_party(row[idx["反担保主体"]])
                if mode == "first":
                    # 一转: 担保方 -> 反担保
                    add_pair(gua, fan, oid)
                else:
                    # 二转: 反担保 -> 受让方
                    add_pair(fan, to, oid)
        wb.close()
        return ordered, groups

    # Mode B: flat case sheet
    ws = wb[sheet_names[0]]
    rows = ws.iter_rows(values_only=True)
    headers = next(rows, None)
    if not headers:
        wb.close()
        raise ValueError("输入表无表头")
    idx = _header_index(list(headers))
    for required in ("订单号", "受让方名称", "担保方名称", "反担保主体"):
        if required not in idx:
            raise ValueError(f"缺少列: {required}")
    has_error = "_主表错误" in idx

    for row in rows:
        if not row or row[idx["订单号"]] in (None, ""):
            continue
        oid = row[idx["订单号"]]
        gua = clean_party(row[idx["担保方名称"]])
        to = clean_party(row[idx["受让方名称"]])
        fan = clean_party(row[idx["反担保主体"]])
        err = str(row[idx["_主表错误"]] or "") if has_error else ""

        # Always send current-holder notice: 反担保 -> 受让方
        add_pair(fan, to, oid)

        # Dual-notice cases: also 担保方 -> 反担保
        need_first = False
        if has_error and "0/2" in err:
            need_first = True
        elif (not has_error) and fan and gua and fan != gua:
            # Fallback when no material-error column: different counter-guarantor
            need_first = True
        if need_first:
            add_pair(gua, fan, oid)

    wb.close()
    return ordered, groups


def write_sms_file(
    template: Path,
    out_path: Path,
    from_party: str,
    to_party: str,
    order_ids: list,
    fund_name: str = DEFAULT_FUND,
) -> int:
    shutil.copy2(template, out_path)
    wb = load_workbook(out_path)
    ws = wb.active

    # Drop legacy send-time column if present
    for col in range(1, ws.max_column + 1):
        if ws.cell(1, col).value == "短信发送时间":
            ws.delete_cols(col)
            break
    if "AF" in ws.column_dimensions:
        del ws.column_dimensions["AF"]

    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)

    for i, oid in enumerate(order_ids):
        row = i + 2
        ws.cell(row, COL_ORDER_ID).value = oid
        ws.cell(row, COL_CURRENT_HOLDER).value = from_party
        ws.cell(row, COL_SECOND_TRANSFEREE).value = from_party
        ws.cell(row, COL_THIRD_TRANSFEROR).value = to_party
        ws.cell(row, COL_FUND).value = fund_name

    wb.save(out_path)
    wb.close()
    return len(order_ids)


def generate(
    input_path: Path,
    out_dir: Path,
    date_mmdd: str,
    template: Path | None = None,
    fund_name: str = DEFAULT_FUND,
    dry_run: bool = False,
    skip_excel_saveas: bool = False,
) -> list[tuple[str, int]]:
    if not re_fullmatch_mmdd(date_mmdd):
        raise ValueError(f"--date 应为 MMDD，例如 0911，收到: {date_mmdd!r}")

    ordered, groups = read_transfer_groups(input_path)
    if not ordered:
        raise ValueError("未解析到任何转让组合")

    plan: list[tuple[str, str, str, list]] = []
    for from_party, to_party in ordered:
        oids = _dedupe(groups[(from_party, to_party)])
        fname = output_filename(from_party, to_party, date_mmdd)
        plan.append((fname, from_party, to_party, oids))

    print(f"input: {input_path}")
    print(f"groups: {len(plan)}")
    for fname, from_party, to_party, oids in plan:
        print(f"  {len(oids):5d}  {short_of(from_party)}转{short_of(to_party)}  -> {fname}")

    if dry_run:
        return [(fname, len(oids)) for fname, _, _, oids in plan]

    out_dir.mkdir(parents=True, exist_ok=True)
    if template is None:
        template = find_default_template([out_dir, input_path.parent, Path.cwd()])
    if template is None or not template.exists():
        raise FileNotFoundError(
            "未找到 FJ-债转短信 模板。请用 --template 指定一份已能上传法久的 xlsx。"
        )
    print(f"template: {template}")

    created: list[tuple[str, int]] = []
    for fname, from_party, to_party, oids in plan:
        out_path = out_dir / fname
        n = write_sms_file(template, out_path, from_party, to_party, oids, fund_name)
        created.append((fname, n))
        print(f"wrote {fname} ({n})")

    if not skip_excel_saveas:
        print("Excel SaveAs for format check...")
        for fname, _ in created:
            excel_saveas(out_dir / fname)
            print(f"excel-saved {fname} size={ (out_dir / fname).stat().st_size }")

    print(f"TOTAL files={len(created)} rows={sum(n for _, n in created)}")
    return created


def re_fullmatch_mmdd(value: str) -> bool:
    return len(value) == 4 and value.isdigit()


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="生成法久债转短信上传表")
    parser.add_argument("--input", "-i", required=True, help="源 Excel（债转短信* 或 全部订单/短信*）")
    parser.add_argument("--out-dir", "-o", required=True, help="输出目录")
    parser.add_argument("--date", "-d", required=True, help="文件名日期 MMDD，如 0911")
    parser.add_argument("--template", "-t", default=None, help="FJ-债转短信模板 xlsx（建议已能上传法久）")
    parser.add_argument("--fund", default=DEFAULT_FUND, help=f"资金方名称，默认 {DEFAULT_FUND}")
    parser.add_argument("--dry-run", action="store_true", help="只打印分组，不写文件")
    parser.add_argument(
        "--skip-excel-saveas",
        action="store_true",
        help="跳过 Excel COM 另存（本地无 Excel 时用）",
    )
    return parser


def main(argv: list[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    input_path = Path(args.input).expanduser().resolve()
    out_dir = Path(args.out_dir).expanduser().resolve()
    template = Path(args.template).expanduser().resolve() if args.template else None
    if not input_path.exists():
        raise SystemExit(f"输入文件不存在: {input_path}")
    generate(
        input_path=input_path,
        out_dir=out_dir,
        date_mmdd=args.date,
        template=template,
        fund_name=args.fund,
        dry_run=args.dry_run,
        skip_excel_saveas=args.skip_excel_saveas,
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
