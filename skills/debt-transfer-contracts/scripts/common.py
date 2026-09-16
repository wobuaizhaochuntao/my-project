# Shared helpers for Fajiu debt-transfer contract generation.
from __future__ import annotations

import re
from copy import deepcopy
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from typing import Any

import openpyxl
from docx import Document
from docx.oxml.ns import qn

Q = Decimal("0.01")
CN_NUM = "零壹贰叁肆伍陆柒捌玖"


def s(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def money(v: Any) -> Decimal:
    try:
        return Decimal(s(v) or "0").quantize(Q, rounding=ROUND_HALF_UP)
    except Exception:
        return Decimal("NaN")


def strip_account_suffix(name: str) -> str:
    return re.sub(r"\d{4}$", "", s(name))


def cn_money(value: Any) -> str:
    d = money(value)
    integer = int(d)
    jiao = int(d * 10) % 10
    fen = int(d * 100) % 10
    units = ["", "拾", "佰", "仟"]
    groups = ["", "万", "亿", "兆"]
    if integer == 0:
        result = "零"
    else:
        parts = []
        gi = 0
        need_zero = False
        while integer:
            g = integer % 10000
            integer //= 10000
            if g == 0:
                if parts:
                    need_zero = True
            else:
                chunk = ""
                zero = False
                for i in range(4):
                    n = (g // (10**i)) % 10
                    if n:
                        if zero and chunk and not chunk.startswith("零"):
                            chunk = "零" + chunk
                        chunk = CN_NUM[n] + units[i] + chunk
                        zero = False
                    elif chunk:
                        zero = True
                if need_zero or (integer and g < 1000):
                    chunk = "零" + chunk
                parts.insert(0, chunk + groups[gi])
                need_zero = False
            gi += 1
        result = re.sub("零+", "零", "".join(parts)).rstrip("零")
    result += "元"
    if jiao == 0 and fen == 0:
        return result + "整"
    if jiao:
        result += CN_NUM[jiao] + "角"
    elif fen:
        result += "零"
    if fen:
        result += CN_NUM[fen] + "分"
    return result


def read_debt_workbook(path: str | Path) -> dict[str, list[list[str]]]:
    """Return {sheet_name: [[序号,订单号,...], ...]} with blank rows skipped."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    data: dict[str, list[list[str]]] = {}
    for ws in wb.worksheets:
        ws.reset_dimensions()
        it = ws.iter_rows(values_only=True)
        headers = [s(x) for x in next(it)]
        rows = []
        for r in it:
            vals = [s(x) for x in r[:8]]
            if any(vals):
                rows.append(vals)
        data[ws.title] = rows
    wb.close()
    return data


def read_unpushed_orders(path: str | Path) -> dict[str, dict[str, Any]]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    ws.reset_dimensions()
    it = ws.iter_rows(values_only=True)
    headers = [s(x) for x in next(it)]
    ix = {h: i for i, h in enumerate(headers)}
    orders: dict[str, dict[str, Any]] = {}
    for r in it:
        oid = s(r[ix["订单号"]])
        if not oid:
            continue
        orders[oid] = {h: r[i] for h, i in ix.items()}
    wb.close()
    return orders


def expected_sheets_for_order(
    order: dict[str, Any],
    full_to_short: dict[str, str],
) -> set[str]:
    guarantor = s(order.get("担保方名称"))
    counter = strip_account_suffix(s(order.get("反担保主体")))
    spv = s(order.get("受让方名称"))
    g = full_to_short[guarantor]
    c = full_to_short[counter]
    sp = full_to_short[spv]
    if guarantor == counter:
        return {f"{g}_{sp}"}
    return {f"{g}_{c}", f"{c}_{sp}"}


def visible_match(masked: Any, full: Any) -> bool:
    a, b = s(masked), s(full)
    return len(a) == len(b) and all(x == "*" or x == y for x, y in zip(a, b))


def source_amounts(order: dict[str, Any]) -> tuple[Decimal, Decimal, Decimal]:
    principal = (money(order.get("代偿本金")) + money(order.get("回购本金"))).quantize(Q)
    interest = (
        money(order.get("代偿利息(含罚息)"))
        + money(order.get("回购利息"))
        + money(order.get("回购罚息"))
        + money(order.get("回购其他"))
    ).quantize(Q)
    return principal, interest, (principal + interest).quantize(Q)


def find_docx(folder: str | Path, kind: str) -> Path:
    folder = Path(folder)
    matches = [p for p in folder.glob("*.docx") if kind in p.name and not p.name.startswith("~$")]
    if len(matches) != 1:
        raise RuntimeError(f"{folder} 中 {kind} 模板数量异常: {[p.name for p in matches]}")
    return matches[0]


def all_paragraphs(doc: Document):
    for p in doc.paragraphs:
        yield p
    for section in doc.sections:
        for p in section.header.paragraphs:
            yield p
        for p in section.footer.paragraphs:
            yield p
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    yield p


def replace_paragraph(p, replacements, regexes=()):
    for run in p.runs:
        t = run.text
        for a, b in replacements:
            t = t.replace(a, b)
        run.text = t
    full = "".join(r.text for r in p.runs)
    new = full
    for a, b in replacements:
        new = new.replace(a, b)
    for pat, repl in regexes:
        new = re.sub(pat, repl, new)
    if new != full:
        if p.runs:
            p.runs[0].text = new
            for r in p.runs[1:]:
                r.text = ""
        else:
            p.add_run(new)


def replace_doc_text(doc: Document, replacements, regexes=()):
    for p in all_paragraphs(doc):
        replace_paragraph(p, replacements, regexes)


def detail_table(doc: Document):
    for t in doc.tables:
        if t.rows and t.rows[0].cells and t.rows[0].cells[0].text.strip() == "序号":
            return t
    return None


def fill_detail_table(doc: Document, rows: list[list[str]], six_cols: bool = False) -> None:
    target = detail_table(doc)
    if target is None:
        raise RuntimeError("未找到附件明细表")
    if len(target.rows) < 2:
        raise RuntimeError("模板明细表没有样式行")
    sample = deepcopy(target.rows[1]._tr)
    for tr in list(target._tbl.tr_lst)[1:]:
        target._tbl.remove(tr)
    for source in rows:
        vals = source[:5] + [source[7]] if six_cols else source[:8]
        # normalize amount display to 2 decimals for money columns
        if six_cols:
            vals[5] = f"{money(vals[5]):.2f}"
        else:
            vals[5] = f"{money(vals[5]):.2f}"
            vals[6] = f"{money(vals[6]):.2f}"
            vals[7] = f"{money(vals[7]):.2f}"
        tr = deepcopy(sample)
        cells = tr.findall(qn("w:tc"))
        if len(cells) != len(vals):
            raise RuntimeError(f"列数不匹配 {len(cells)} != {len(vals)}")
        for tc, val in zip(cells, vals):
            texts = list(tc.iter(qn("w:t")))
            if not texts:
                raise RuntimeError("样式单元格缺少文本节点")
            texts[0].text = str(val)
            for node in texts[1:]:
                node.text = ""
        target._tbl.append(tr)


def fmt_cn_date(ymd: str) -> str:
    """2026-09-10 -> 2026年09月10日"""
    y, m, d = ymd.split("-")
    return f"{y}年{m}月{d}日"


def fmt_bracket_date(ymd: str) -> str:
    y, m, d = ymd.split("-")
    return f"【{y}】年【{m}】月【{d}】日"


def ymd_compact(ymd: str) -> str:
    return ymd.replace("-", "")
