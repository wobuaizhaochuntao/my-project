#!/usr/bin/env python3
"""Fill供应商/客户录入明细 workbook for one or more SPVs."""
from __future__ import annotations

import argparse
import json
from copy import copy
from pathlib import Path

from openpyxl import load_workbook


def copy_row_style(ws, source_row: int, target_row: int, max_col: int) -> None:
    ws.row_dimensions[target_row].height = ws.row_dimensions[source_row].height
    for col in range(1, max_col + 1):
        s = ws.cell(source_row, col)
        t = ws.cell(target_row, col)
        if s.has_style:
            t._style = copy(s._style)
        if s.number_format:
            t.number_format = s.number_format
        t.font = copy(s.font)
        t.fill = copy(s.fill)
        t.border = copy(s.border)
        t.alignment = copy(s.alignment)
        t.protection = copy(s.protection)


def main() -> int:
    ap = argparse.ArgumentParser(description="按 JSON 列表写入供应商客户录入明细")
    ap.add_argument("--template", required=True, help="原录入明细 xlsx 模板")
    ap.add_argument("--out", required=True, help="输出 xlsx 路径")
    ap.add_argument("--companies", required=True, help="公司 JSON 数组文件")
    args = ap.parse_args()

    companies = json.loads(Path(args.companies).read_text(encoding="utf-8"))
    wb = load_workbook(args.template)

    contact_name = companies[0].get("contact_name", "赵纯涛")
    contact_phone = companies[0].get("contact_phone", "15907105801")
    contact_dept = companies[0].get("contact_dept", "运营")
    coop_status = companies[0].get("coop_status", "合作中")
    tax_rate = companies[0].get("tax_rate", 0.06)
    our_dept = companies[0].get("our_dept", "运营中心&法务部")

    ws = wb["基本信息"]
    for row in range(3, 2 + len(companies)):
        copy_row_style(ws, 2, row, 15)
    for i, c in enumerate(companies, 2):
        vals = [
            "供应商",
            c["name"],
            "否",
            c.get("address", ""),
            c.get("established", ""),
            c.get("credit", ""),
            c.get("legal", ""),
            c.get("capital", ""),
            c.get("scope", ""),
            contact_name,
            contact_phone,
            contact_dept,
            coop_status,
            tax_rate,
            our_dept,
        ]
        for col, val in enumerate(vals, 1):
            ws.cell(i, col).value = val if val != "" else None

    ws = wb["账户信息"]
    for row in range(3, 2 + len(companies)):
        copy_row_style(ws, 2, row, 4)
    for i, c in enumerate(companies, 2):
        for col, val in enumerate(
            [c["name"], c.get("bank", ""), c.get("account", ""), "人民币"], 1
        ):
            ws.cell(i, col).value = val
    for row in range(2 + len(companies), ws.max_row + 1):
        for col in range(1, 5):
            ws.cell(row, col).value = None

    ws = wb["开票信息"]
    for row in range(3, 2 + len(companies)):
        copy_row_style(ws, 2, row, 5)
    for i, c in enumerate(companies, 2):
        for col, val in enumerate(
            [
                c["name"],
                c.get("address", ""),
                contact_phone,
                c.get("bank", ""),
                c.get("account", ""),
            ],
            1,
        ):
            ws.cell(i, col).value = val

    Path(args.out).parent.mkdir(parents=True, exist_ok=True)
    wb.save(args.out)
    print(args.out)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
