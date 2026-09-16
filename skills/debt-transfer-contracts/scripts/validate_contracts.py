#!/usr/bin/env python3
"""Validate generated contracts against债转文件 and optional未推送订单."""
from __future__ import annotations

import argparse
import json
import re
import sys
import zipfile
from collections import defaultdict
from decimal import Decimal
from pathlib import Path

ROOT = Path(__file__).resolve().parent
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from common import (  # noqa: E402
    detail_table,
    expected_sheets_for_order,
    fmt_cn_date,
    money,
    read_debt_workbook,
    read_unpushed_orders,
    source_amounts,
    visible_match,
    ymd_compact,
)
from docx import Document  # noqa: E402


def xml_text(path: Path) -> str:
    with zipfile.ZipFile(path) as z:
        return "".join(
            z.read(n).decode("utf-8", "ignore")
            for n in z.namelist()
            if n.endswith(".xml")
        )


def table_rows_match(doc: Document, expected: list[list[str]], inds: list[int]) -> str | None:
    t = detail_table(doc)
    if t is None:
        return "缺少明细表"
    if len(t.rows) != len(expected) + 1 or len(t.columns) != len(inds):
        return f"表格尺寸 {len(t.rows)-1}x{len(t.columns)} 预期 {len(expected)}x{len(inds)}"
    for i, (tr, er) in enumerate(zip(t.rows[1:], expected), 1):
        got = [c.text.strip() for c in tr.cells]
        exp = [er[j] for j in inds]
        for gi, (a, b) in enumerate(zip(got, exp)):
            # money columns: compare numerically
            if (len(inds) == 6 and gi == 5) or (len(inds) == 8 and gi in (5, 6, 7)):
                if money(a) != money(b):
                    return f"第{i}行金额不一致"
            elif a != b:
                return f"第{i}行不一致: {a!r} vs {b!r}"
    return None


def audit_routing(debt: dict, orders: dict, full_to_short: dict) -> dict:
    mem = defaultdict(list)
    for sh, rows in debt.items():
        for r in rows:
            mem[r[1]].append(sh)
    logic_bad = []
    for oid, o in orders.items():
        exp = expected_sheets_for_order(o, full_to_short)
        act = set(mem.get(oid, []))
        if act != exp:
            logic_bad.append({"order": oid, "expected": sorted(exp), "actual": sorted(act)})
    field_bad = 0
    for oid, sheets in mem.items():
        if oid not in orders:
            continue
        o = orders[oid]
        principal, interest, total = source_amounts(o)
        for sh in sheets:
            row = next(r for r in debt[sh] if r[1] == oid)
            checks = [
                row[2] == str(o.get("被申请人姓名") or "").strip(),
                visible_match(row[3], o.get("身份证号码")),
                visible_match(row[4], o.get("被申请人联系方式")),
                money(row[5]) == principal,
                money(row[6]) == interest,
                money(row[7]) == total,
            ]
            if not all(checks):
                field_bad += 1
                break
    return {
        "source_orders": len(orders),
        "debt_unique": len(mem),
        "missing_in_debt": sorted(set(orders) - set(mem))[:20],
        "missing_in_source": sorted(set(mem) - set(orders))[:20],
        "logic_mismatches": len(logic_bad),
        "logic_samples": logic_bad[:10],
        "order_field_mismatches": field_bad,
    }


def main() -> int:
    ap = argparse.ArgumentParser(description="核对债转文件 / 合同 / 未推送订单")
    ap.add_argument("--config", required=True, help="与生成时相同的 JSON 配置")
    ap.add_argument("--unpushed", help="未推送订单.xlsx（可选）")
    args = ap.parse_args()

    cfg = json.loads(Path(args.config).read_text(encoding="utf-8"))
    debt = read_debt_workbook(cfg["debt_excel"])
    errors: list[str] = []
    root = Path(cfg["output_root"])
    batch_n = cfg["batch_number"]
    batch_label = cfg["batch_label"]
    first_mmdd = ymd_compact(cfg["dates"]["first"])[4:]
    second_mmdd = ymd_compact(cfg["dates"]["second"])[4:]
    out_first = root / f"法久第{batch_n}批一转协议{first_mmdd}"
    out_second = root / f"法久第{batch_n}批二转协议{second_mmdd}"

    checked_docs = 0
    checked_rows = 0

    for item in cfg.get("first_transfers", []):
        rec = cfg["parties"][item["recipient"]]
        compact = ymd_compact(cfg["dates"]["first"])
        path = out_first / f"【上海法久&{rec['short']}】一转协议-{compact}.docx"
        if not path.exists():
            errors.append(f"缺文件 {path}")
            continue
        doc = Document(str(path))
        err = table_rows_match(doc, debt[item["sheet"]], [0, 1, 2, 3, 4, 7])
        text = "\n".join(p.text for p in doc.paragraphs)
        if err:
            errors.append(f"{path.name}: {err}")
        if rec["name"] not in text and rec["name"] not in xml_text(path):
            errors.append(f"{path.name}: 缺主体 {rec['name']}")
        first_cn = fmt_cn_date(cfg["dates"]["first"])
        if first_cn not in text:
            errors.append(f"{path.name}: 缺日期 {first_cn}")
        checked_docs += 1
        checked_rows += len(debt[item["sheet"]])

    for item in cfg.get("second_groups", []):
        tr = cfg["parties"][item["transferor"]]
        dest = cfg["parties"][item["recipient"]]
        sheet = item["sheet"]
        compact = ymd_compact(cfg["dates"]["second"])
        folder = out_second / sheet
        for kind, suffix in (("二转协议", ""), ("三转协议", "-2"), ("补充协议", "-1")):
            path = folder / f"【法久&{tr['short']}&{batch_label}】{kind}-{compact}.docx"
            if not path.exists():
                errors.append(f"缺文件 {path}")
                continue
            doc = Document(str(path))
            text = "\n".join(p.text for p in doc.paragraphs)
            x = xml_text(path)
            num = f"{item['transferor']}-{item['recipient']}-{compact}{suffix}"
            if num not in text:
                errors.append(f"{path.name}: 缺编号 {num}")
            if f"【{len(debt[sheet])}】笔" not in text:
                errors.append(f"{path.name}: 缺笔数")
            if tr["name"] not in x or dest["name"] not in x:
                errors.append(f"{path.name}: 主体不完整")
            if kind != "补充协议":
                err = table_rows_match(doc, debt[sheet], list(range(8)))
                if err:
                    errors.append(f"{path.name}: {err}")
                else:
                    checked_rows += len(debt[sheet])
            if kind == "三转协议":
                if not any(p.text.strip() == "签订日期：" for p in doc.paragraphs):
                    errors.append(f"{path.name}: 三转签订日期未留空")
            else:
                cn = fmt_cn_date(cfg["dates"]["second"])
                if cn not in text:
                    errors.append(f"{path.name}: 缺日期 {cn}")
            stale = cfg.get("stale_tokens", [])
            for tok in stale:
                if tok in x:
                    errors.append(f"{path.name}: 残留 {tok}")
            checked_docs += 1

    routing = None
    if args.unpushed:
        orders = read_unpushed_orders(args.unpushed)
        full_to_short = {p["name"]: p["short"] for p in cfg["parties"].values()}
        # Also map names that only appear as guarantor aliases if provided
        for alias, short in cfg.get("name_aliases", {}).items():
            full_to_short[alias] = short
        routing = audit_routing(debt, orders, full_to_short)

    result = {
        "checked_docs": checked_docs,
        "checked_attachment_rows": checked_rows,
        "errors": errors,
        "routing": routing,
        "ok": len(errors) == 0 and (routing is None or routing["logic_mismatches"] == 0),
    }
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0 if result["ok"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
